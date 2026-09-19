from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

import pymupdf
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, SimpleTestCase, override_settings
from django.urls import reverse

from .chat_files import export_answer, prepare_input
from .models import AIProfile, ModelSetting, UsageLog


class FileConversionTests(SimpleTestCase):
    def test_document_roundtrips(self):
        text = '日本語の回答\n=1+1\n次の行'
        doc = export_answer(text, 'docx')
        self.assertEqual('\n'.join(p.text for p in Document(BytesIO(doc)).paragraphs), text)
        xlsx = export_answer(text, 'xlsx')
        book = load_workbook(BytesIO(xlsx))
        self.assertEqual(book.active['A2'].value, '=1+1')
        self.assertEqual(book.active['A2'].data_type, 's')
        for name, raw in [('sample.docx', doc), ('sample.xlsx', xlsx)]:
            content, audit = prepare_input('', [SimpleUploadedFile(name, raw)], 1000)
            self.assertIn('日本語の回答', content[0]['content'][0]['text'])
            self.assertIn(name, audit)

    def test_pdf_and_images(self):
        for format in ('pdf', 'jpeg', 'png'):
            data = export_answer('日本語の回答', format)
            if format == 'pdf':
                with pymupdf.open(stream=data, filetype='pdf') as pdf:
                    self.assertIn('日本語の回答', pdf[0].get_text())
            else:
                with Image.open(BytesIO(data)) as image:
                    self.assertEqual(image.format, format.upper())
            content, _ = prepare_input('', [SimpleUploadedFile('sample.' + format, data)], 1000)
            self.assertEqual(content[0]['content'][1]['type'], 'input_file' if format == 'pdf' else 'input_image')

    def test_long_images_keep_all_pages(self):
        data = export_answer('回答\n' * 100, 'png')
        with ZipFile(BytesIO(data)) as archive:
            self.assertEqual(len(archive.namelist()), 3)

    def test_reject_invalid_uploads_and_limits(self):
        for filename, data in [('x.exe', b'x'), ('x.png', b'fake'), ('x.docx', b'fake'), ('x.pdf', b'fake'), ('x.png', b'')]:
            with self.subTest(filename=filename), self.assertRaises(ValueError):
                prepare_input('', [SimpleUploadedFile(filename, data)], 1000)
        with self.assertRaises(ValueError):
            prepare_input('', [], 1000)
        with self.assertRaises(ValueError):
            prepare_input('1234', [], 3)
        with self.assertRaises(ValueError):
            prepare_input('', [SimpleUploadedFile('x.docx', export_answer('long text', 'docx'))], 3)
        with self.assertRaises(ValueError):
            prepare_input('', [SimpleUploadedFile('x.png', b'x') for _ in range(6)], 1000)


@override_settings(OPENAI_API_KEY='test-key')
class ChatAttachmentTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('attachment-user')
        AIProfile.objects.create(user=self.user)
        self.model = ModelSetting.objects.create(model_code='attachment-model', display_name='Test')
        self.client.force_login(self.user)

    @patch('openai.OpenAI')
    def test_upload_download_and_ownership(self, api):
        api.return_value.responses.create.return_value = SimpleNamespace(
            output_text='日本語の回答', usage=SimpleNamespace(input_tokens=12, output_tokens=10))
        response = self.client.post(reverse('internal_ai:chat_api'), {
            'model': self.model.model_code,
            'files': SimpleUploadedFile('input.docx', export_answer('質問の本文', 'docx')),
        })
        self.assertEqual(response.status_code, 200, response.content)
        self.assertIn('質問の本文', api.return_value.responses.create.call_args.kwargs['input'][0]['content'][0]['text'])
        self.assertEqual(UsageLog.objects.get(user=self.user).status, 'SUCCESS')
        for url in response.json()['downloads'].values():
            result = self.client.get(url)
            self.assertEqual(result.status_code, 200)
            self.assertIn('attachment;', result['Content-Disposition'])
        url = response.json()['downloads']['pdf']
        other = get_user_model().objects.create_user('other')
        self.client.force_login(other)
        self.assertEqual(self.client.get(url).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

    @patch('openai.OpenAI')
    def test_invalid_file_does_not_call_api(self, api):
        response = self.client.post(reverse('internal_ai:chat_api'), {
            'model': self.model.model_code, 'files': SimpleUploadedFile('x.png', b'fake')})
        self.assertEqual(response.status_code, 400)
        api.assert_not_called()

    def test_previews_paginate_and_require_ownership(self):
        from uuid import uuid4
        log = UsageLog.objects.create(request_id=uuid4(), user=self.user, model=self.model,
                                      status='SUCCESS', response_text='日本語の回答\n' * 100)
        for format in ('png', 'jpeg', 'pdf'):
            url = reverse('internal_ai:chat_download', args=[log.request_id, format])
            result = self.client.get(url, {'preview': '1', 'page': '2'})
            self.assertEqual(result.status_code, 200)
            self.assertTrue(result['Content-Disposition'].startswith('inline;'))
            self.assertEqual(result['Cache-Control'], 'private, no-store')
            if format == 'pdf':
                with pymupdf.open(stream=result.content, filetype='pdf') as pdf:
                    self.assertEqual(len(pdf), 3)
            else:
                self.assertEqual(result['X-Page-Count'], '3')
                with Image.open(BytesIO(result.content)) as image:
                    self.assertEqual(image.format, format.upper())
                for page, status in [('bad', 400), ('0', 404), ('4', 404)]:
                    self.assertEqual(self.client.get(url, {'preview': '1', 'page': page}).status_code, status)
        other = get_user_model().objects.create_user('preview-other')
        self.client.force_login(other)
        self.assertEqual(self.client.get(url, {'preview': '1'}).status_code, 404)
        self.client.force_login(self.user)
        AIProfile.objects.filter(user=self.user).update(enabled=False)
        self.assertEqual(self.client.get(url, {'preview': '1'}).status_code, 403)

    def test_preview_rejects_office_format(self):
        from uuid import uuid4
        log = UsageLog.objects.create(request_id=uuid4(), user=self.user, model=self.model,
                                      status='SUCCESS', response_text='回答')
        url = reverse('internal_ai:chat_download', args=[log.request_id, 'docx'])
        self.assertEqual(self.client.get(url, {'preview': '1'}).status_code, 400)

    @patch('openai.OpenAI')
    def test_generate_real_image_and_download(self, api):
        import base64
        buffer = BytesIO()
        Image.new('RGB', (32, 32), 'red').save(buffer, 'PNG')
        raw = buffer.getvalue()
        api.return_value.responses.create.return_value = SimpleNamespace(
            output_text='', usage=SimpleNamespace(input_tokens=12, output_tokens=10),
            output=[SimpleNamespace(type='image_generation_call', result=base64.b64encode(raw).decode())])
        response = self.client.post(reverse('internal_ai:chat_api'), {
            'model': self.model.model_code, 'message': '赤い四角の画像を作って', 'answer_type': 'image'})
        self.assertEqual(response.status_code, 200, response.content)
        options = api.return_value.responses.create.call_args.kwargs
        self.assertEqual(options['tool_choice'], {'type': 'image_generation'})
        self.assertEqual(options['tools'][0]['output_format'], 'png')
        self.assertEqual(api.call_args.kwargs['timeout'], 180)
        data = response.json()
        self.assertTrue(data['image_cost_estimated'])
        self.assertEqual(len(data['images']), 1)
        log = UsageLog.objects.get(user=self.user)
        self.assertTrue(log.image_cost_estimated)
        self.assertGreaterEqual(log.cost_usd, Decimal('0.20'))
        for format, url in data['images'][0].items():
            preview = self.client.get(url)
            self.assertEqual(preview.status_code, 200)
            self.assertTrue(preview['Content-Disposition'].startswith('inline;'))
            self.assertEqual(preview['Cache-Control'], 'private, no-store')
            with Image.open(BytesIO(preview.content)) as image:
                self.assertEqual(image.size, (32, 32))
                self.assertGreater(image.getpixel((0, 0))[0], 250)
            if format == 'png':
                self.assertEqual(preview.content, raw)
            self.assertTrue(self.client.get(url, {'download': '1'})['Content-Disposition'].startswith('attachment;'))
        self.client.force_login(get_user_model().objects.create_user('image-other'))
        self.assertEqual(self.client.get(url).status_code, 404)
        self.client.force_login(self.user)
        AIProfile.objects.filter(user=self.user).update(enabled=False)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

    @patch('openai.OpenAI')
    def test_image_refusal_and_text_mode(self, api):
        api.return_value.responses.create.return_value = SimpleNamespace(
            output_text='この画像は生成できません。', output=[],
            usage=SimpleNamespace(input_tokens=10, output_tokens=5))
        for mode in ('image', 'text'):
            response = self.client.post(reverse('internal_ai:chat_api'), {
                'model': self.model.model_code, 'message': '質問', 'answer_type': mode})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()['images'], [])
            self.assertEqual(response.json()['answer'], 'この画像は生成できません。')
            self.assertEqual('tools' in api.return_value.responses.create.call_args.kwargs, mode == 'image')

    @patch('openai.OpenAI')
    def test_image_mode_obeys_limits(self, api):
        AIProfile.objects.filter(user=self.user).update(daily_request_limit=0)
        response = self.client.post(reverse('internal_ai:chat_api'), {
            'message': '画像を作って', 'answer_type': 'image'})
        self.assertEqual(response.status_code, 403)
        api.assert_not_called()
