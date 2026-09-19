"""Bounded attachment parsing and downloadable answer documents."""
import base64
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
import pymupdf

FORMATS = {
    'txt': 'text/plain; charset=utf-8',
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pdf': 'application/pdf', 'jpeg': 'image/jpeg', 'png': 'image/png',
}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_FILES = 5


def prepare_input(message, files, max_chars):
    if len(files) > MAX_FILES or sum(f.size for f in files) > 20 * 1024 * 1024:
        raise ValueError('添付は5件、合計20MBまでです。')
    if not message.strip() and not files:
        raise ValueError('質問またはファイルを入力してください。')
    content = []
    texts = [message] if message.strip() else []
    for upload in files:
        ext = Path(upload.name).suffix.lower()
        if ext not in ('.docx', '.xlsx', '.pdf', '.jpg', '.jpeg', '.png'):
            raise ValueError('添付形式は.docx、.xlsx、.pdf、.jpg、.jpeg、.pngに対応しています。')
        if not upload.size or upload.size > MAX_FILE_BYTES:
            raise ValueError('各ファイルは空でない10MB以下のファイルにしてください。')
        raw = upload.read(MAX_FILE_BYTES + 1)
        try:
            if ext in ('.docx', '.xlsx'):
                with ZipFile(BytesIO(raw)) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                        raise ValueError('展開後のファイルが大きすぎます。')
                parts = []
                count = 0
                def add(text):
                    nonlocal count
                    count += len(text) + 1
                    if count > max_chars:
                        raise ValueError('添付内の文字数が入力上限を超えています。')
                    parts.append(text)
                if ext == '.docx':
                    doc = Document(BytesIO(raw))
                    for item in doc.iter_inner_content():
                        if hasattr(item, 'rows'):
                            for row in item.rows:
                                add('\t'.join(cell.text for cell in row.cells))
                        else:
                            add(item.text)
                else:
                    book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
                    try:
                        cells = 0
                        for sheet in book:
                            add('シート: ' + sheet.title)
                            for row in sheet.iter_rows():
                                cells += len(row)
                                if cells > 100000:
                                    raise ValueError('Excelのセル数が多すぎます。')
                                add('\t'.join('' if c.value is None else str(c.value) for c in row))
                    finally:
                        book.close()
                texts.append(f'添付: {upload.name}\n' + '\n'.join(parts))
            elif ext == '.pdf':
                with pymupdf.open(stream=raw, filetype='pdf') as pdf:
                    if pdf.needs_pass or not 0 < len(pdf) <= 100:
                        raise ValueError('PDFは暗号化されていない100ページ以下のものにしてください。')
                content.append({'type': 'input_file', 'filename': upload.name,
                                'file_data': 'data:application/pdf;base64,' + base64.b64encode(raw).decode()})
            else:
                with Image.open(BytesIO(raw)) as img:
                    expected = 'PNG' if ext == '.png' else 'JPEG'
                    if img.format != expected or img.width * img.height > 20000000:
                        raise ValueError('画像形式または画像サイズが不正です。')
                    img.verify()
                mime = 'image/png' if ext == '.png' else 'image/jpeg'
                content.append({'type': 'input_image', 'detail': 'auto',
                                'image_url': f'data:{mime};base64,' + base64.b64encode(raw).decode()})
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f'「{upload.name}」を読み取れません。ファイルの破損や暗号化を確認してください。') from exc
    text = '\n\n'.join(texts)
    if len(text) > max_chars:
        raise ValueError('質問と添付内の文字数の合計が入力上限を超えています。')
    if not files:
        return message, message
    content.insert(0, {'type': 'input_text', 'text': text or '添付ファイルの内容を説明してください。'})
    audit = text + '\n' + '\n'.join('添付ファイル: ' + f.name for f in files)
    return [{'role': 'user', 'content': content}], audit


def export_answer(text, format):
    buffer = BytesIO()
    if format == 'txt':
        return text.encode('utf-8')
    if format == 'docx':
        doc = Document()
        for line in text.split('\n'):
            doc.add_paragraph(line)
        doc.save(buffer)
    elif format == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = 'AI回答'
        sheet.column_dimensions['A'].width = 100
        from openpyxl.styles import Alignment
        for line in text.split('\n'):
            for start in range(0, max(1, len(line)), 32000):
                sheet.append([line[start:start + 32000]])
                cell = sheet.cell(sheet.max_row, 1)
                cell.data_type = 's'  # Never execute generated text as spreadsheet formulas.
                cell.alignment = Alignment(wrap_text=True, vertical='top')
        book.save(buffer)
    elif format in ('pdf', 'jpeg', 'png'):
        # MuPDF includes a Japanese font, independent of host font installation.
        font = pymupdf.Font('japan')
        lines = []
        for paragraph in text.split('\n'):
            line = ''
            width = 0
            for char in paragraph.expandtabs(4):
                advance = font.text_length(char, fontsize=11)
                if width + advance > 495 and line:
                    lines.append(line)
                    line, width = '', 0
                line += char
                width += advance
            lines.append(line)
        with pymupdf.open() as pdf:
            for start in range(0, len(lines), 44):
                page = pdf.new_page(width=595, height=842)
                page.insert_font(fontname='answer', fontbuffer=font.buffer)
                page.insert_text((50, 60), '\n'.join(lines[start:start + 44]), fontname='answer', fontsize=11, lineheight=1.5)
            if format == 'pdf':
                return pdf.tobytes(garbage=4, deflate=True)
            if len(pdf) == 1:
                return pdf[0].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).tobytes(format)
            # Preserve every page without creating an unbounded tall bitmap.
            with ZipFile(buffer, 'w') as archive:
                for index, page in enumerate(pdf):
                    archive.writestr(f'answer-{index + 1}.{format}', page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5)).tobytes(format))
    else:
        raise ValueError('対応していない出力形式です。')
    return buffer.getvalue()


def validate_generated_image(encoded):
    if not isinstance(encoded, str) or len(encoded) > 30 * 1024 * 1024:
        raise ValueError('生成画像のサイズが不正です。')
    raw = base64.b64decode(encoded, validate=True)
    with Image.open(BytesIO(raw)) as image:
        if image.format != 'PNG' or image.width * image.height > 20000000:
            raise ValueError('生成画像の形式が不正です。')
        image.verify()
